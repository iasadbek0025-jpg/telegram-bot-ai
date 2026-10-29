"""Tadbirkorlik boshqarmasi uchun Telegram yordamchi bot (mavzular + hujjat tahlili)."""
import asyncio
import io
import logging
import os
import re
 
import docx
import pypdf
from aiogram import Bot, Dispatcher, F
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import Command, CommandObject, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import KeyboardButton, Message, ReplyKeyboardMarkup
from dotenv import load_dotenv
from google import genai
from google.genai import types
 
load_dotenv()
 
# ---------------------------------------------------------------- sozlamalar
BOT_TOKEN = os.getenv("BOT_TOKEN", "")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
# Ixtiyoriy: ruxsat berilgan foydalanuvchi ID'lari (vergul bilan). Bo'sh bo'lsa - hamma uchun.
ALLOWED_USERS = {int(x) for x in os.getenv("ALLOWED_USERS", "").split(",") if x.strip()}
 
MAX_FILE_MB = 15
MAX_CHARS = 60_000          # Gemini'ga yuboriladigan matn chegarasi
TG_LIMIT = 4000             # Telegram xabari 4096 belgidan oshmasligi kerak
 
BTN_TOPICS = "📊 Mavzular yaratish"
BTN_FILE = "📁 Faylni tahlil qilish"
 
DISCLAIMER = (
    "\n\n⚠️ Bu sun'iy intellekt tayyorlagan dastlabki material. "
    "Rasmiy foydalanishdan oldin raqam, sana va me'yoriy hujjatlarni mutaxassis tekshirishi shart."
)
 
TOPICS_SYSTEM = (
    "Siz tadbirkorlikni qo'llab-quvvatlash boshqarmasining iqtisodiy tahlilchisi va matbuot "
    "xizmati mutaxassisisiz. O'zbek tilida (lotin yozuvida), sodda va rasmiy uslubda yozing.\n"
    "Qoidalar: raqam, foiz, sana, qonun/qaror raqamlarini O'YLAB TOPMANG. Aniq ma'lumot "
    "ishonchli bo'lmasa, o'rniga [manba va raqam tekshirilsin] deb yozing."
)
TOPICS_TASK = (
    "Quyidagi yo'nalish bo'yicha tadbirkorlar uchun foydali 5 ta Telegram post mavzusini tuzing.\n"
    "Har bir mavzu uchun: 1) sarlavha, 2) 3-4 bandli qisqa reja, "
    "3) post uchun tekshirilishi kerak bo'lgan manba (qaysi qonun yoki idora ma'lumoti).\n"
    "Yo'nalish: {direction}"
)
 
ANALYSIS_SYSTEM = (
    "Siz tadbirkorlikni qo'llab-quvvatlash boshqarmasining bosh tahlilchisisiz. "
    "O'zbek tilida (lotin yozuvida), rasmiy va aniq uslubda yozing.\n"
    "Qoidalar: faqat hujjatdagi ma'lumotga tayaning; hujjatda yo'q raqam, modda yoki "
    "me'yoriy aktni o'ylab topmang; taxminni 'taxmin' deb belgilang; yuridik xulosa o'rniga "
    "'yurist tekshiruvi tavsiya etiladi' deb yozing."
)
ANALYSIS_TASK = (
    "Hujjatni tekshirib, quyidagi tuzilmada xulosa bering:\n"
    "1. Qisqa mazmun (3-5 jumla)\n"
    "2. Asosiy qoida va raqamlar\n"
    "3. Kamchiliklar va noaniq joylar\n"
    "4. Tadbirkorlar uchun muhim yoki xavfli jihatlar\n"
    "5. Takomillashtirish bo'yicha tavsiyalar"
)
 
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("bot")
 
bot = Bot(BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
dp = Dispatcher()
client = genai.Client(api_key=GEMINI_API_KEY)
busy_users: set[int] = set()
 
 
class Form(StatesGroup):
    topic = State()
 
 
# ---------------------------------------------------------------- yordamchilar
def main_keyboard() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text=BTN_TOPICS), KeyboardButton(text=BTN_FILE)]],
        resize_keyboard=True,
    )
 
 
def allowed(message: Message) -> bool:
    return not ALLOWED_USERS or (message.from_user and message.from_user.id in ALLOWED_USERS)
 
 
def clean_markdown(text: str) -> str:
    """Gemini Markdown'ini Telegram uchun oddiy matnga aylantiradi."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"^\s*#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"^\s*[\*\-]\s+", "• ", text, flags=re.M)
    return text.replace("`", "").strip()
 
 
def split_text(text: str, limit: int = TG_LIMIT) -> list[str]:
    parts, current = [], ""
    for line in text.split("\n"):
        while len(line) > limit:                       # juda uzun qator
            if current:
                parts.append(current)
                current = ""
            parts.append(line[:limit])
            line = line[limit:]
        if len(current) + len(line) + 1 > limit:
            parts.append(current)
            current = line
        else:
            current = f"{current}\n{line}" if current else line
    if current:
        parts.append(current)
    return parts
 
 
async def send_long(message: Message, text: str) -> None:
    for chunk in split_text(text):
        await message.answer(chunk, parse_mode=None)
 
 
async def ask_gemini(contents, system: str, retries: int = 2) -> str:
    config = types.GenerateContentConfig(system_instruction=system, temperature=0.3)
    for attempt in range(retries + 1):
        try:
            resp = await client.aio.models.generate_content(
                model=GEMINI_MODEL, contents=contents, config=config
            )
            text = (resp.text or "").strip()
            if text:
                return clean_markdown(text)
            raise RuntimeError("Gemini bo'sh javob qaytardi")
        except Exception:
            log.exception("Gemini xatosi (urinish %s)", attempt + 1)
            if attempt == retries:
                raise
            await asyncio.sleep(2 * (attempt + 1))
 
 
# ---------------------------------------------------------------- fayllardan matn olish
def read_txt(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")
 
 
def read_docx(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:                              # jadvallar ham o'qiladi
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)
 
 
def read_pdf(data: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)
 
 
# ---------------------------------------------------------------- handlerlar
@dp.message(CommandStart())
async def start_handler(message: Message, state: FSMContext):
    await state.clear()
    if not allowed(message):
        await message.answer("Kechirasiz, sizda bu botdan foydalanish huquqi yo'q.")
        return
    await message.answer(
        "Assalomu alaykum! 🏛 <b>Tadbirkorlik boshqarmasi</b> raqamli yordamchisiga xush kelibsiz.\n\n"
        "• Tadbirkorlar uchun post mavzularini tayyorlayman\n"
        "• <b>PDF, Word (.docx), TXT</b> hujjatlarni tahlil qilaman\n\n"
        "<i>Javoblar sun'iy intellekt tomonidan tayyorlanadi va mutaxassis tekshiruvini talab qiladi.</i>",
        reply_markup=main_keyboard(),
    )
 
 
@dp.message(F.text == BTN_TOPICS)
async def topics_button(message: Message, state: FSMContext):
    if not allowed(message):
        return
    await state.set_state(Form.topic)
    await message.answer("Qaysi yo'nalish bo'yicha mavzu kerak? Yo'nalishni yozing.\nMasalan: <i>Soliq imtiyozlari</i>")
 
 
@dp.message(F.text == BTN_FILE)
async def file_button(message: Message, state: FSMContext):
    if not allowed(message):
        return
    await state.clear()
    await message.answer(f"Hujjatni <b>PDF, DOCX yoki TXT</b> formatida yuboring (hajmi {MAX_FILE_MB} MB gacha).")
 
 
async def generate_topics(message: Message, direction: str):
    user_id = message.from_user.id
    if user_id in busy_users:
        await message.answer("Oldingi so'rovingiz bajarilmoqda, iltimos kuting.")
        return
    busy_users.add(user_id)
    try:
        await message.answer("⏳ Mavzular tayyorlanmoqda...")
        result = await ask_gemini(TOPICS_TASK.format(direction=direction), TOPICS_SYSTEM)
        await send_long(message, f"📋 Mavzular loyihasi:\n\n{result}{DISCLAIMER}")
    except Exception:
        await message.answer("Xizmatda vaqtincha xatolik yuz berdi. Birozdan so'ng qayta urinib ko'ring.")
    finally:
        busy_users.discard(user_id)
 
 
@dp.message(Command("mavzu"))
async def mavzu_command(message: Message, command: CommandObject, state: FSMContext):
    if not allowed(message):
        return
    await state.clear()
    direction = (command.args or "").strip()
    if not direction:
        await state.set_state(Form.topic)
        await message.answer("Yo'nalishni yozing. Masalan: <i>Subsidiyalar va imtiyozli kreditlar</i>")
        return
    await generate_topics(message, direction)
 
 
@dp.message(Form.topic, F.text)
async def topic_received(message: Message, state: FSMContext):
    if not allowed(message):
        return
    direction = message.text.strip()
    if len(direction) < 3 or len(direction) > 300:
        await message.answer("Yo'nalish 3 dan 300 belgigacha bo'lishi kerak. Qayta yozing.")
        return
    await state.clear()
    await generate_topics(message, direction)
 
 
@dp.message(F.document)
async def handle_document(message: Message, state: FSMContext):
    if not allowed(message):
        return
    await state.clear()
    doc = message.document
    name = (doc.file_name or "").lower()
 
    if not name.endswith((".txt", ".docx", ".pdf")):
        await message.answer("⚠️ Faqat .pdf, .docx yoki .txt fayllar qabul qilinadi.")
        return
    if doc.file_size and doc.file_size > MAX_FILE_MB * 1024 * 1024:
        await message.answer(f"⚠️ Fayl hajmi {MAX_FILE_MB} MB dan oshmasligi kerak.")
        return
 
    user_id = message.from_user.id
    if user_id in busy_users:
        await message.answer("Oldingi so'rovingiz bajarilmoqda, iltimos kuting.")
        return
    busy_users.add(user_id)
 
    try:
        await message.answer("🔍 Fayl qabul qilindi, tahlil qilinmoqda... ⏳")
        buffer = io.BytesIO()
        await bot.download(doc, destination=buffer)
        data = buffer.getvalue()
 
        scanned_pdf = False
        if name.endswith(".txt"):
            text = read_txt(data)
        elif name.endswith(".docx"):
            text = await asyncio.to_thread(read_docx, data)
        else:
            text = await asyncio.to_thread(read_pdf, data)
            scanned_pdf = len(text.strip()) < 50       # skaner qilingan PDF
 
        note = ""
        if scanned_pdf:
            # Matn yo'q: PDF'ni Gemini'ning o'zi ko'rib o'qiydi
            contents = [types.Part.from_bytes(data=data, mime_type="application/pdf"), ANALYSIS_TASK]
            note = "ℹ️ PDF skaner qilingan ko'rinadi, shuning uchun rasm sifatida o'qildi.\n\n"
        else:
            if not text.strip():
                await message.answer("❌ Faylda o'qiladigan matn topilmadi.")
                return
            if len(text) > MAX_CHARS:
                text = text[:MAX_CHARS]
                note = f"ℹ️ Hujjat uzun, faqat birinchi {MAX_CHARS:,} belgisi tahlil qilindi.\n\n"
            contents = f"{ANALYSIS_TASK}\n\nHUJJAT MATNI:\n{text}"
 
        result = await ask_gemini(contents, ANALYSIS_SYSTEM)
        await send_long(message, f"📊 Hujjat tahlili:\n\n{note}{result}{DISCLAIMER}")
    except Exception:
        log.exception("Hujjatni qayta ishlashda xatolik")
        await message.answer("Faylni qayta ishlashda xatolik yuz berdi. Fayl buzilmaganini tekshirib, qayta yuboring.")
    finally:
        busy_users.discard(user_id)
 
 
async def main():
    log.info("Bot ishga tushdi. Model: %s", GEMINI_MODEL)
    await dp.start_polling(bot)
 
 
if __name__ == "__main__":
    asyncio.run(main())
 